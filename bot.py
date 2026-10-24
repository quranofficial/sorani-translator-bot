#!/usr/bin/env python3
# -*- coding: utf-8 -*-
# =====================================================================
#  بۆتی وەرگێڕی کوردی سۆرانی  |  Kurdish Sorani Translator Bot
# =====================================================================
#  تواناکان / Features:
#   - دەق (text) ............ وەرگێڕان ڕاستەوخۆ بۆ کوردی سۆرانی
#   - PDF / DOCX / PPTX / TXT / SRT ... وەرگێڕانی فایل، دەرچووەکە هەمان جۆر
#        + کۆپی PDF بۆ هەموو فایلەکان (DOCX/PPTX بە LibreOffice بۆ باوەڕپێکراوی شێواز)
#   - ڤیدیۆ / دەنگ (فایل یان لینکی YouTube / Facebook / TikTok / ...)
#        دەنگ -> نووسین (faster-whisper) -> وەرگێڕان -> SRT + Word + PDF بە سۆرانی
#        + ژێرنووسی کوردی ڕاستەوخۆ لەسەر ڤیدیۆکە دادەنرێت (burn-in بە ffmpeg/libass)
#   - بزوێنەری وەرگێڕان: Google Translate (بەخۆڕایی) + GLM/Z.ai (یەدەگ)
#   - لە GitHub Actions ٢٤/٧ کار دەکات: هەر خولێک ٤٥ خولەکە و لە کۆتایی خولێکی نوێ
#     خۆکارانە دەست پێدەکات (زنجیرەی بێکۆتایی) + cron ی پاشەکەوت هەر ١٠ خولەک
#
#  کاردانەوەی خێرا: python bot.py --check   (پشکنینی ژینگە)
# =====================================================================

import glob
import io
import json
import logging
import os
import random
import re
import shutil
import subprocess
import sys
import tempfile
import time
import traceback
from concurrent.futures import ThreadPoolExecutor

import requests

# ------------------------------ Config -------------------------------

def env_int(name, default):
    v = (os.getenv(name) or "").strip()
    try:
        return int(float(v)) if v else default
    except ValueError:
        return default

BOT_TOKEN         = (os.getenv("BOT_TOKEN") or "").strip()
TARGET_LANG       = (os.getenv("TARGET_LANG") or "ckb").strip()      # ckb = کوردی سۆرانی
STATE_FILE        = (os.getenv("STATE_FILE") or "state.json").strip()
# خاڵی Telegram: ئەگەر سێرڤەری ناوخۆیی Bot API کاردەکات (docker)، فایل تا 2GB دابەردەکرێت
TG_BASE           = (os.getenv("TG_API_BASE") or "https://api.telegram.org").strip().rstrip("/")
TG_FILE_BASE      = (os.getenv("TG_FILE_BASE") or TG_BASE).strip().rstrip("/")
LOCAL_BOT_API     = ("localhost" in TG_BASE) or ("127.0.0.1" in TG_BASE)
RUN_SECONDS       = env_int("RUN_SECONDS", 2700)     # ماوەی خولی polling (بنەڕەت ٤٥ خولەک)
IDLE_EXIT_SECONDS = env_int("IDLE_EXIT_SECONDS", 0)  # 0 = لەکارخراو — بۆت تا کۆتایی خول ئۆنلاین دەمێنێت
HEAVY_BUDGET      = env_int("HEAVY_BUDGET", 2400)    # زۆرترین کات بۆ یەک ڤیدیۆ/دەنگ (چرکە)
MAX_VIDEO_MINUTES = env_int("MAX_VIDEO_MINUTES", 0)  # 0 = بێ سنوور
BURN_SUBTITLES    = env_int("BURN_SUBTITLES", 1)     # 1 = ژێرنووسی کوردی لەسەر ڤیدیۆ دادەنرێت
BURN_MAX_MINUTES  = env_int("BURN_MAX_MINUTES", 30)  # زۆرترین درێژی ڤیدیۆ بۆ سووتاندنی ژێرنووس
SEND_PDF_COPY     = env_int("SEND_PDF_COPY", 1)      # 1 = کۆپی PDF بۆ هەموو فایلەکان دەنێردرێت
OUT_VIDEO_MAX_MB  = env_int("OUT_VIDEO_MAX_MB", 20)  # ڤیدیۆی دەرچوو خۆکارانە بۆ خوار ئەم قەبارە بچووک دەکرێتەوە
SHRINK_OVER_MB    = env_int("SHRINK_OVER_MB", 60)    # ڤیدیۆی گەورەتر لەمە پێش پرۆسێسکردن بۆ 720p بچووک دەکرێتەوە
WHISPER_MODEL     = (os.getenv("WHISPER_MODEL") or "base").strip()
SOURCE_LANG       = (os.getenv("SOURCE_LANG") or "").strip() or None  # None = خۆکار
GLM_API_KEY       = (os.getenv("GLM_API_KEY") or os.getenv("ZAI_API_KEY") or "").strip()
GLM_API_URL       = (os.getenv("GLM_API_URL") or "https://api.z.ai/api/paas/v4").strip().rstrip("/")
# زنجیرەی مۆدێلەکانی GLM: یەکەم دانەی کارا لەبەر دەکرێت (glm-4.5-flash خۆڕاییە)
_glm_model_env    = (os.getenv("GLM_MODEL") or "").strip()
GLM_MODELS        = ([_glm_model_env] if _glm_model_env else []) + [
    m.strip() for m in (os.getenv("GLM_MODELS") or
                        "glm-4.5-flash,glm-4-flash,glm-4.5-air").split(",") if m.strip()]
_glm_seen = set()
GLM_MODELS = [m for m in GLM_MODELS if m and not (m in _glm_seen or _glm_seen.add(m))]
YT_COOKIES        = (os.getenv("YT_COOKIES") or "").strip()
# --- دەنگی کوردی سۆرانی (TTS) — kurdishtts.com ---
TTS_API_KEY   = (os.getenv("KURDISH_TTS_API_KEY") or os.getenv("TTS_API_KEY") or "").strip()
TTS_API_URL   = (os.getenv("KURDISH_TTS_API_URL") or "https://www.kurdishtts.com/api/tts-proxy").strip()
TTS_SPEAKER   = (os.getenv("TTS_SPEAKER_ID") or "sorani_85").strip()   # خۆڕایی: sorani_85 / sorani_214
TTS_MODEL_V   = (os.getenv("TTS_MODEL_VERSION") or "v3").strip()
TTS_SPEED     = float((os.getenv("TTS_SPEED") or "1.0").strip() or "1.0")
TTS_MAX_CHARS = env_int("TTS_MAX_CHARS", 15000)  # سەقەتی پیت بۆ هەر ڤیدیۆیەک (پلانی خۆڕایی: 20000/مانگ)

TG_API = TG_BASE + "/bot" + BOT_TOKEN
TG_MAX_FILE = (2000 if LOCAL_BOT_API else 20) * 1024 * 1024

logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")
log = logging.getLogger("sorani-bot")

SESSION = requests.Session()
SESSION.headers.update({"User-Agent": "Mozilla/5.0 (X11; Linux x86_64) sorani-translator-bot/1.0"})


class BotError(Exception):
    """هەڵەی دۆستانە: نامەکەی وەکو خۆی بۆ بەکارهێنەر دەنێردرێت."""


# -------------------------- نووسینە کوردییەکان ------------------------

WELCOME = (
    "سڵاو! من بۆتی وەرگێڕانم بۆ کوردی سۆرانی.\n"
    "\n"
    "چی دەکەم:\n"
    "• هەر دەقێک بنێری (هەر زمانێک بێت) — ڕاستەوخۆ بە سۆرانی وەریدەگێڕمەوە.\n"
    "• فایلی PDF، Word، PowerPoint، TXT یان SRT بنێری — وەرگێڕدراوەکە بە هەمان جۆری فایل\n"
    "  دەنێرمەوە + کۆپی PDF یش لەگەڵ دەکەم.\n"
    "• ڤیدیۆ یان دەنگ بنێری (فایل)، یان لینکی YouTube / Facebook / TikTok — دەنگەکە دەخوێنمەوە،\n"
    "  وەریدەگێڕم و بە فایلی SRT و Word و PDF دەینێرم + ژێرنووسی کوردی لەسەر ڤیدیۆکە دادەنێم.\n"
    "• ئەگەر ڤیدیۆکە لە ٢٠ مێگابایت گەورەتر بوو، بۆ خوار ٢٠ مێگابایت بچووکی دەکەمەوە.\n"
    "• دەنگی کوردی سۆرانی لەسەر ڤیدیۆ (ئارەزوومەندانە): /voice_on بۆ چالاککردن،\n"
    "  /voice_off بۆ ناچالاککردن — /voice بۆ زانینی دۆخ.\n"
    "\n"
    "تێبینی گرنگ:\n"
    "• فایلی ڕاستەوخۆ دەبێت کەمتر بێت لە ٢٠ مێگابایت (سنووری تەلەگرام).\n"
    "• بەردەوام ئۆنلاینم (٢٤ کاتژمێر / ٧ ڕۆژ)؛ تەنها لە کاتی گۆڕینی خولەکە (هەر نزیکەی کاتژمێرێک)\n"
    "  لەوانەیە وەڵامەکەم چەند خولەکێک دابخرێت."
)

PROG_DOWNLOAD   = "چاوەڕێ بکە... ڤیدیۆکە داگیر دەم. (بەپێی قەبارەکەی چەند خولەکێک دەخایەنێت)"
MSG_MEDIA_START = "چاوەڕێ بکە... فایلەکە دادەگرم و دەنگەکە دەخوێنمەوە."
MSG_DOC_START   = "چاوەڕێ بکە... فایلەکە وەرگێڕ دەم. (بۆ فایلە گەورەکان چەند خولەکێک دەخایەنێت)"
MSG_DOWNLOADED  = "داگیرا: {title}\nماوە: {dur} — ئێستا دەنگ دەخوێنمەوە..."
MSG_TRANSCRIBED = "دەنگ خرایە سەر نووسین: زمانی «{lang}» — ماوە {dur} — {n} بەش.\nئێستا وەرگێڕان بۆ سۆرانی دەست پێدەکات..."
MSG_VIDEO_DONE  = "وەرگێڕان تەواو بوو: فایلی SRT بۆ ژێرنووس + فایلی Word و PDF بۆ خوێندنەوە."
CAPTION_DOC     = "وەرگێڕانی کوردی سۆرانی ئامادەیە."
CAPTION_PDF     = "هەمان وەرگێڕان بە شێوەی PDF."
MSG_BURN_START   = "ئێستا ژێرنووسی کوردی لەسەر ڤیدیۆکە دادەنێم... (ئەم قۆناغە چەند خولەکێک دەخایەنێت)"
CAPTION_BURN     = "ڤیدیۆکە بە ژێرنووسی کوردی سۆرانی."
MSG_BURN_FAIL    = ("بەداخەوە دانانی ژێرنووس لەسەر ڤیدیۆ سەرنەچوو؛ "
                    "بەڵام فایلەکانی SRT و Word و PDF ئامادەن.")
MSG_BURN_SKIPPED = ("ڤیدیۆکە لە {max} خولەک درێژترە، بۆیە ژێرنووسەکە نەخرایە سەر ڤیدیۆ "
                    "(سنووری کاتی GitHub Actions).\n"
                    "فایلی SRT ئامادەیە — دەتوانیت بە بەرنامەی مۆنتاژ یان خزمەتگوزاری "
                    "کاپشن بیخەیتە سەر ڤیدیۆکە.")
MSG_TRANS_PARTIAL = "تێبینی: {n} بەش وەرگێڕانەکەیان سەرنەچوو و بە زمانی ڕەسەن ماونەتەوە."
MSG_COMPRESS = ("قەبارەی ڤیدیۆکە لە {mb} مێگابایت زیاترە — بۆ ناردنەوە بچووکی دەکەمەوە "
                "بۆ خوار {mb} مێگابایت... (چەند خولەکێک دەخایەنێت)")
MSG_ORIG_TRANSCRIPT = "نووسینی ڕەسەن (بەبێ وەرگێڕان) — دووبارە هەوڵ بدە بۆ وەرگێڕان."
MSG_VOICE_ON = "دەنگی کوردی سۆرانی چالاک کرا — لە ڤیدیۆی داهاتوودا دەنگە کوردییەکە لەسەر ڤیدیۆ دادەنێم."
MSG_VOICE_OFF = "دەنگی کوردی ناچالاک کرا — ڤیدیۆکان بە دەنگی ڕەسەنی خۆیان دەنێردرێنەوە."
MSG_VOICE_STATUS_ON = "ئێستا: دەنگی کوردی چالاکە (/voice_off بۆ ناچالاککردن)"
MSG_VOICE_STATUS_OFF = "ئێستا: دەنگی کوردی ناچالاکە (/voice_on بۆ چالاککردن)"
MSG_TTS_START = "ئێستا دەنگی کوردی سۆرانی بۆ ڤیدیۆکە دروست دەکەم... (چەند خولەکێک دەخایەنێت)"
MSG_TTS_FAIL = "دروستکردنی دەنگی کوردی سەرنەچوو — ڤیدیۆکە بە دەنگی ڕەسەن نێردرایەوە."
MSG_TTS_NOKEY = ("بۆ دەنگی کوردی سۆرانی: کلیدێکی TTS لە kurdishtts.com/settings/api بەدەست بهێنە\n"
                 "و لە GitHub بیکە بە secret بە ناوی KURDISH_TTS_API_KEY.\n"
                 "(پلانی خۆڕایی: ٢٠٠٠٠ پیت بۆ مانگێک — دەنگی بنەڕەت: sorani_85)")
MSG_UNSUPPORTED = ("ئەم جۆرە پەیام/فایلە پشتگیری نەکراوە.\n"
                   "پشتگیریکراو: PDF، Word (docx)، PowerPoint (pptx)، TXT، SRT، "
                   "ڤیدیۆ، دەنگ، لینکی ڤیدیۆ، یان دەقی ئاسایی.")
MSG_ALREADY     = "ئەم دەقە پێشتر سۆرانییە یان پێویستی بە وەرگێڕان نییە."
ERR_LINK        = ("نەتوانم ڤیدیۆکە لەم لینکە دابگرم. لەوانەیە ماڵپەڕەکە ڕێگری لە سێرڤەرەکان بکات "
                   "(بەتایبەتی YouTube لە ئایپییەکانی GitHub).\n"
                   "چارەسەر: دووبارە هەوڵ بدە، یان ڤیدیۆکە دابگرە و وەکو فایل بنێرە (تا ٢٠ مێگابایت).")
def err_big():
    """پەیامی فایلی گەورە — بەپێی جۆری خاڵی Telegram."""
    if LOCAL_BOT_API:
        return "فایلەکە لە ٢ گێگابایت گەورەترە — ئەمە زۆرترین سنووری تەلەگرامە بۆ فایل."
    return ("فایلەکە لە ٢٠ مێگابایت گەورەترە — سنووری داگرتنەوەی تەلەگرامە بۆ بۆتەکان.\n"
            "چارەسەر:\n"
            "• لە GitHub دوو secret زیاد بکە: TG_API_ID و TG_API_HASH (لە my.telegram.org بەخۆڕایین)\n"
            "  — ئەوسا فایلەکان تا ٢ گێگابایت دەگرم و خۆکارانە بچووکی دەکەمەوە بۆ خوار ٢٠ مێگابایت.\n"
            "• یان لینکی ڤیدیۆکە بنێرە (YouTube / TikTok / Facebook...).\n"
            "• یان ڤیدیۆکە بچووکتر بکە و دووبارە بنێرە.")
ERR_SCANNED     = ("ئەم فایلی PDF ە دەقی نووسراوی تێدا نییە (وەکو وێنەی سکانیکراوە).\n"
                   "هەوڵی OCR درا بەڵام دەقی گونجاوی لێ دەرنەچوو — لەوانەیە کوالیتی وێنەکان زۆر لاواز بێت.\n"
                   "چارەسەر: وێنەی ڕوونتر بە PDF بنێرە، یان دەقەکە بە TXT بنێرە.")
ERR_LONG        = "پرۆسێسەکە زۆر درێژ کێشا و ڕاوەستێنرا. تکایە ڤیدیۆ/فایلەکە بکە بە بەشی بچووکتر."
ERR_AUDIO       = "نەتوانم دەنگێک لەم فایلە بخوێنمەوە. لەوانەیە بێ دەنگ بێت یان فۆرماتەکە پشتگیری نەکرێت."
ERR_FFMPEG      = "ئامرازی ffmpeg لەم ژینگەیەدا بەردەست نییە (لە GitHub Actions خۆکارانە دامەزرێت)."
ERR_TRANS       = ("وەرگێڕان سەرنەچوو — گووگڵ و GLM هەردووکیان وەڵام نەدانەوە.\n"
                   "بە زۆری دووبارە هەوڵدان سەرکەوتوو دەبێت: تکایە پەیامەکە دووبارە بنێرە.\n"
                   "(ئەگەر بەردەوام بوو: لە GitHub چێک بکە کە secret ی ZAI_API_KEY ڕاستە)")
ERR_GEN         = "هەڵەیەکی نەخوازراو ڕوویدا. تکایە دووبارە هەوڵ بدە."

GLM_PROMPT = (
    "You are a professional translator specialized in Central Kurdish "
    "(Sorani, ISO 639-3: ckb, written in Arabic script). Translate the user's "
    "text into natural, fluent Sorani Kurdish. Preserve the original paragraph "
    "structure and line breaks. Output ONLY the translated text — no "
    "explanations, no quotes, no extra commentary."
)

DOC_EXTS  = {".pdf", ".docx", ".pptx", ".txt", ".srt"}
MEDIA_EXTS = {".mp4", ".mkv", ".webm", ".mov", ".avi", ".m4v", ".3gp", ".ts",
              ".mp3", ".m4a", ".aac", ".wav", ".ogg", ".oga", ".opus", ".flac",
              ".wma", ".amr"}


def has_letters(t):
    return any(c.isalpha() for c in (t or ""))


def sanitize(name):
    name = re.sub(r"[^\w.\- ]+", "_", (name or "file"))[:80].strip()
    return name or "file"


def fmt_dur(sec):
    try:
        sec = int(float(sec))
    except Exception:
        return "?"
    return "%d خولەک و %d چرکە" % (sec // 60, sec % 60)


def ts(sec):
    """چرکە -> 00:00:00,000  (فۆرماتی SRT)"""
    sec = max(0.0, float(sec))
    ms = int(round(sec * 1000))
    h, rem = divmod(ms, 3_600_000)
    m, rem = divmod(rem, 60_000)
    s, ms3 = divmod(rem, 1000)
    return "%02d:%02d:%02d,%03d" % (h, m, s, ms3)


# ----------------------------- Telegram ------------------------------

def tg_call(method, payload=None, files=None, timeout=70):
    r = SESSION.post(TG_API + "/" + method, data=payload, files=files, timeout=timeout)
    try:
        j = r.json()
    except Exception:
        raise RuntimeError("telegram %s http %s" % (method, r.status_code))
    if not j.get("ok"):
        desc = j.get("description") or ("http %s" % r.status_code)
        if "too big" in desc.lower():
            raise BotError(err_big())
        raise BotError("Telegram: " + desc)
    return j.get("result")


def tg_send(chat_id, text):
    text = (text or "").strip() or "-"
    while text:
        if len(text) <= 4000:
            part, rest = text, ""
        else:
            cut = text.rfind("\n", 0, 4000)
            if cut < 2000:
                cut = 4000
            part, rest = text[:cut], text[cut:].lstrip("\n")
        tg_call("sendMessage", {"chat_id": chat_id, "text": part})
        text = rest


def _safe_send(chat_id, text):
    try:
        tg_send(chat_id, text)
    except Exception:
        log.error("failed to send message to %s", chat_id)


def tg_action(chat_id, action):
    try:
        tg_call("sendChatAction", {"chat_id": chat_id, "action": action}, timeout=20)
    except Exception:
        pass


def tg_send_doc(chat_id, path, caption=None):
    payload = {"chat_id": chat_id}
    if caption:
        payload["caption"] = caption[:1000]
    with open(path, "rb") as fh:
        tg_call("sendDocument", payload,
                files={"document": (os.path.basename(path), fh)}, timeout=600)


def tg_download(file_id, workdir, name_hint="file"):
    f = tg_call("getFile", {"file_id": file_id}, timeout=60)
    if int(f.get("file_size") or 0) > TG_MAX_FILE:
        raise BotError(err_big())
    url = "%s/file/bot%s/%s" % (TG_FILE_BASE, BOT_TOKEN, f["file_path"])
    r = SESSION.get(url, timeout=600)
    r.raise_for_status()
    ext = os.path.splitext(f.get("file_path") or "")[1] or ""
    path = os.path.join(workdir, sanitize(name_hint) + ext)
    with open(path, "wb") as fh:
        fh.write(r.content)
    return path


def tg_get_updates(offset, timeout=40):
    try:
        return tg_call("getUpdates",
                       {"offset": offset, "timeout": int(timeout), "limit": 100,
                        "allowed_updates": ["message"]},
                       timeout=int(timeout) + 20) or []
    except requests.RequestException as e:
        log.warning("network error while polling: %s", e)
        return []
    except BotError as e:
        if "409" in str(e) or "conflict" in str(e).lower():
            log.warning("getUpdates conflict (409) — another poller is running; sleeping")
            time.sleep(20)
            return []
        raise


def skip_backlog(state):
    """لە یەکەم کارکردن: تەنها کۆتا نوێکردنەوە وەربگرە، پەیامە کۆنەکان تێپەڕە."""
    res = tg_call("getUpdates", {"offset": -1, "limit": 1, "timeout": 0}, timeout=30)
    if res:
        state["offset"] = res[-1]["update_id"] + 1
        save_state(state)
        log.info("backlog skipped, starting offset=%s", state["offset"])


def load_state():
    try:
        with open(STATE_FILE, "r", encoding="utf-8") as fh:
            return json.load(fh)
    except Exception:
        return {"offset": 0}


def save_state(state):
    tmp = STATE_FILE + ".tmp"
    with open(tmp, "w", encoding="utf-8") as fh:
        json.dump(state, fh)
    os.replace(tmp, STATE_FILE)


# -------------------------- وەرگێڕان ---------------------------------
#  Google Translate — چەند ڕێڕەو (translate.googleapis.com، clients5، translate.google.com،
#  هەرکام بە POST و GET) + GLM/Z.ai یەدەگ بە چەند مۆدێل و چەند سێرڤەر.
#  ئەمە وا دەکات وەرگێڕان لە GitHub Actions دا بەردەوام بێ (429 ی گووگڵ و مۆدێلی هەڵەی GLM چارە دەکات).

GTX_HOSTS = [
    # (ناونیشان, client, dj) — dj=1 → وەڵامی JSON ی نوێ؛ dict-chrome-ex → array ی کۆن
    ("https://translate.googleapis.com/translate_a/single", "gtx", True),
    ("https://clients5.google.com/translate_a/single", "dict-chrome-ex", False),
    ("https://translate.google.com/translate_a/single", "gtx", True),
]
GTX_UA = ("Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 "
          "(KHTML, like Gecko) Chrome/124.0 Safari/537.36")
GTX_ROUNDS = env_int("GTX_ROUNDS", 2)
_MEMO = {}


def _parse_gtx(data, text):
    """پارسکردنی وەڵامی گووگڵ — هەردوو فۆرمات (dj=1 ی dict و array ی کۆن)."""
    out, src = "", ""
    if isinstance(data, dict):
        src = data.get("src") or ""
        out = "".join(s.get("trans", "") for s in (data.get("sentences") or []))
    elif isinstance(data, list) and data and isinstance(data[0], list):
        for seg in (data[0] or []):
            if seg and isinstance(seg[0], str):
                out += seg[0]
        if len(data) > 2 and isinstance(data[2], str):
            src = data[2]
    if not out.strip() and text.strip():
        raise RuntimeError("gtx empty result")
    return out, src


def _gtx_request(host, client, use_dj, text, use_get):
    params = {"client": client, "sl": SOURCE_LANG or "auto", "tl": TARGET_LANG, "dt": "t"}
    if use_dj:
        params["dj"] = "1"
    hdrs = {"User-Agent": GTX_UA, "Accept": "*/*"}
    if use_get:
        params["q"] = text
        r = SESSION.get(host, params=params, headers=hdrs, timeout=30)
    else:
        r = SESSION.post(host, params=params, data={"q": text}, headers=hdrs, timeout=30)
    if r.status_code != 200:
        raise RuntimeError("gtx http %s" % r.status_code)
    ctype = (r.headers.get("content-type") or "").lower()
    if not any(x in ctype for x in ("json", "javascript", "text/plain")):
        raise RuntimeError("gtx bad content-type: %s" % ctype[:50])
    return _parse_gtx(r.json(), text)


def _gtx_once(text):
    """سازگاری — یەک هەوڵ بە یەکەم ڕێڕەو (بۆ --check و تاقیکردنەوە)."""
    host, client, use_dj = GTX_HOSTS[0]
    return _gtx_request(host, client, use_dj, text, use_get=False)


def _gtx_translate(text):
    """هەوڵ بە هەموو ڕێڕەوەکانی گووگڵ (POST پاشان GET، دوو خول) + پشووی بەربڵاو."""
    last, attempt = None, 0
    for _rnd in range(GTX_ROUNDS):
        for host, client, use_dj in GTX_HOSTS:
            for use_get in (False, True):
                attempt += 1
                try:
                    out, src = _gtx_request(host, client, use_dj, text, use_get)
                    if src == TARGET_LANG:      # پێشتر سۆرانییە
                        return text
                    return out
                except Exception as e:
                    last = e
                    time.sleep(min(0.8 + 1.1 * attempt + random.random() * 1.5, 7))
    raise last or RuntimeError("gtx failed")


GLM_ALT_URLS = ["https://api.z.ai/api/paas/v4", "https://open.bigmodel.cn/api/paas/v4"]
_glm_combos_cache = None
_glm_good = None          # (url, model) ی ئەو تێکەڵەی کار دەکات — لەبەر دەکرێت


def _glm_combos():
    """هەموو تێکەڵەکانی (سێرڤەر، مۆدێل) — بە ڕیزبەندی هەوڵ دەدرێن."""
    global _glm_combos_cache
    if _glm_combos_cache is None:
        urls = [GLM_API_URL] + [u for u in GLM_ALT_URLS if u.rstrip("/") != GLM_API_URL]
        _glm_combos_cache = [(u, m) for u in urls for m in GLM_MODELS]
    return _glm_combos_cache


def _glm_call(url, model, text):
    r = SESSION.post(
        url + "/chat/completions",
        headers={"Authorization": "Bearer " + GLM_API_KEY},
        json={
            "model": model,
            "temperature": 0.2,
            "messages": [
                {"role": "system", "content": GLM_PROMPT},
                {"role": "user", "content": text},
            ],
        },
        timeout=180,
    )
    if r.status_code in (401, 403):
        raise RuntimeError("glm %s http %s @ %s (key not accepted)"
                           % (model, r.status_code, url))
    if r.status_code != 200:
        raise RuntimeError("glm %s http %s: %s" % (model, r.status_code, (r.text or "")[:150]))
    j = r.json()
    out = ((j.get("choices") or [{}])[0].get("message") or {}).get("content") or ""
    out = re.sub(r"(?s)<think>.*?</think>", "", out).strip()
    if not out:
        raise RuntimeError("glm %s empty result" % model)
    return out


def _glm_once(text):
    """GLM بە چەند (سێرڤەر، مۆدێل) — تێکەڵەی سەرکەوتوو بۆ خولانی داهاتوو لەبەر دەکرێت."""
    global _glm_good
    if not GLM_API_KEY:
        raise RuntimeError("GLM not configured (set ZAI_API_KEY secret)")
    combos = _glm_combos()
    if _glm_good and _glm_good in combos:
        combos = [_glm_good] + [c for c in combos if c != _glm_good]
    last = None
    for url, model in combos:
        try:
            out = _glm_call(url, model, text)
            if _glm_good != (url, model):
                log.info("GLM working combo: %s / %s", url, model)
            _glm_good = (url, model)
            return out
        except Exception as e:
            last = e
            log.warning("%s", e)
    raise last or RuntimeError("glm failed")


def chunk_text(text, size=2800):
    """دەبەشکات بۆ پارچەی <=size، بەبێ لەدەستدانی هیچ پیتێک (بەپێی دێر)."""
    if len(text) <= size:
        return [text]
    lines = text.splitlines(keepends=True)
    chunks, cur = [], ""
    for ln in lines:
        while len(ln) > size:
            if cur:
                chunks.append(cur)
                cur = ""
            chunks.append(ln[:size])
            ln = ln[size:]
        if len(cur) + len(ln) > size:
            chunks.append(cur)
            cur = ln
        else:
            cur += ln
    if cur:
        chunks.append(cur)
    return [c for c in chunks if c]


def _translate_with_fallback(chunk):
    try:
        return _gtx_translate(chunk)
    except Exception as ge:
        log.warning("all Google routes failed (%s), trying GLM fallback", ge)
    try:
        return _glm_once(chunk)
    except Exception as e:
        log.error("GLM fallback failed: %s", e)
        raise BotError(ERR_TRANS)


def translate_text(text):
    text = text or ""
    if not has_letters(text):
        return text
    if len(text) <= 2800:
        return _translate_with_fallback(text)
    chunks = chunk_text(text)
    outs, nfail = translate_map_safe(chunks)
    if nfail >= len(chunks):
        raise BotError(ERR_TRANS)      # هەموو پارچەکان سەرنەچوون
    return "".join(outs)


def translate_item(text):
    """وەرگێڕانی پاراگراف/بەشێک بە یادەوەری (memo) — بۆ فایلەکان."""
    key = (text or "").strip()
    if not key:
        return text
    if key in _MEMO:
        return _MEMO[key]
    out = translate_text(text)
    _MEMO[key] = out
    return out


def translate_map_safe(texts, workers=6):
    """وەرگێڕانی لیستێک بە پاراستنی بەردەوامی: ئەگەر پارچەیەک سەرنەچوو،
    دەقی ڕەسەن دەمێنێتەوە و تەنها ژمارەکەی دەگەڕێنرێتەوە.
    دەرەنجام: (لیستی ئەنجامەکان, ژمارەی پارچە سەرنەچووەکان)"""
    texts = list(texts)
    if not texts:
        return [], 0

    def _one(t):
        try:
            return translate_item(t), True
        except Exception as e:
            log.error("chunk translation failed (%s) — keeping original: %.60r", e, t)
            return t, False

    with ThreadPoolExecutor(max(1, min(workers, len(texts)))) as ex:
        res = list(ex.map(_one, texts))
    return [r[0] for r in res], sum(1 for _, ok in res if not ok)


# -------------------------- دۆکیومێنتەکان -----------------------------

# --- TXT ---
def process_txt(data):
    text = data.decode("utf-8", errors="replace")
    return translate_text(text).encode("utf-8")


# --- SRT (ژێرنووس — کاتەکان دەمێننەوە، تەنها دەق وەرگێڕدرێت) ---
SRT_BLOCK_RE = re.compile(
    r"(\d+)\s*\n"
    r"(\d{1,2}:\d{2}:\d{2}[,.]\d{1,3}\s*-->\s*\d{1,2}:\d{2}:\d{2}[,.]\d{1,3})[^\n]*\n"
    r"(.*?)(?=\n\s*\n|\Z)",
    re.S,
)


def parse_srt(data):
    text = data.decode("utf-8", errors="replace").replace("\r\n", "\n").replace("\r", "\n")
    cues = []
    for m in SRT_BLOCK_RE.finditer(text):
        stamp = m.group(2).strip()
        stamp = stamp.replace(".", ",") if "-->" in stamp else stamp
        body = m.group(3).strip()
        if body:
            cues.append((stamp, body))
    return cues


def translate_srt(data):
    """(srt_وەرگێڕدراو, cues, outs) دەگەڕێنێتەوە — بۆ SRT و PDF ی لگەڵ."""
    cues = parse_srt(data)
    if not cues:
        return "", [], []
    outs, nfail = translate_map_safe([c[1] for c in cues])
    if nfail >= len(cues):
        raise BotError(ERR_TRANS)
    parts = []
    for i, ((stamp, _), out) in enumerate(zip(cues, outs), 1):
        parts.append("%d\n%s\n%s\n" % (i, stamp, (out or "").strip()))
    return "\n".join(parts), cues, outs


def process_srt(data):
    return translate_srt(data)[0]


# --- DOCX (Word) — گۆڕینی دەقی پاراگرافەکان لە شوێنی خۆیان (شێواز دەمێنێتەوە) ---
try:
    from docx import Document as DocxDocument
    from docx.oxml.ns import qn as dx_qn
    try:
        from docx.oxml import OxmlElement
    except Exception:
        from docx.oxml.shared import OxmlElement
except Exception as e:      # pragma: no cover
    raise SystemExit("python-docx is required: pip install python-docx (%s)" % e)


def _set_bidi_docx(p):
    try:
        pPr = p._p.get_or_add_pPr()
        if pPr.find(dx_qn("w:bidi")) is None:
            pPr.append(OxmlElement("w:bidi"))
    except Exception:
        pass


def _set_para_text(p, new_text):
    for hl in p._p.findall(dx_qn("w:hyperlink")):
        p._p.remove(hl)
    if p.runs:
        p.runs[0].text = new_text
        for r in p.runs[1:]:
            r.text = ""
    else:
        p.add_run(new_text)
    _set_bidi_docx(p)


def _docx_paragraphs(doc):
    for p in doc.paragraphs:
        yield p
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    yield p
    for sec in doc.sections:
        for p in sec.header.paragraphs:
            yield p
        for p in sec.footer.paragraphs:
            yield p


def process_docx(data):
    doc = DocxDocument(io.BytesIO(data))
    paras = [p for p in _docx_paragraphs(doc) if has_letters(p.text)]
    if not paras:
        return data
    texts = [p.text for p in paras]
    outs, nfail = translate_map_safe(texts)
    if nfail >= len(texts):
        raise BotError(ERR_TRANS)      # هەموو پاراگرافەکان سەرنەچوون
    for p, out in zip(paras, outs):
        _set_para_text(p, out)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# --- PPTX (PowerPoint) — گۆڕینی دەق لە شوێنی خۆی (لایەوت دەمێنێتەوە) ---
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


def _iter_shapes(shapes):
    for sh in shapes:
        try:
            if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
                for sub in _iter_shapes(sh.shapes):
                    yield sub
                continue
        except Exception:
            pass
        yield sh


def _ppt_paras(tf):
    try:
        return list(tf.paragraphs)
    except Exception:
        return []


def _set_ppt_para(para, new_text):
    try:
        runs = para.runs
        if runs:
            runs[0].text = new_text
            for r in runs[1:]:
                r.text = ""
        else:
            para.add_run().text = new_text
        pPr = para._p.get_or_add_pPr()
        pPr.set("rtl", "1")
    except Exception:
        pass


def process_pptx(data):
    prs = Presentation(io.BytesIO(data))
    items = []
    for slide in prs.slides:
        for sh in _iter_shapes(slide.shapes):
            try:
                if sh.has_text_frame:
                    for para in _ppt_paras(sh.text_frame):
                        if has_letters(para.text):
                            items.append((para, para.text))
            except Exception:
                pass
            try:
                if getattr(sh, "has_table", False) and sh.has_table:
                    for row in sh.table.rows:
                        for cell in row.cells:
                            for para in _ppt_paras(cell.text_frame):
                                if has_letters(para.text):
                                    items.append((para, para.text))
            except Exception:
                pass
        try:
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
                for para in _ppt_paras(slide.notes_slide.notes_text_frame):
                    if has_letters(para.text):
                        items.append((para, para.text))
        except Exception:
            pass
    if not items:
        return data
    texts = [t for _, t in items]
    outs, nfail = translate_map_safe(texts)
    if nfail >= len(texts):
        raise BotError(ERR_TRANS)      # هەموو دەقەکان سەرنەچوون
    for (para, _), out in zip(items, outs):
        _set_ppt_para(para, out)
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


# --- PDF: دەرهێنانی دەق (PyMuPDF) + بنیاتنانی PDF ی سۆرانی RTL (ReportLab) ---
try:
    import pymupdf as fitz_mod
except Exception:
    import fitz as fitz_mod

import arabic_reshaper
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas as rl_canvas

try:
    from bidi.algorithm import get_display
except Exception:
    from bidi import get_display

_reshaper = arabic_reshaper.ArabicReshaper()


def _shape(s):
    return _reshaper.reshape(s)


def _disp(s):
    return get_display(s)


# فۆنتەکان: سەرەتا فۆنتی نێو کۆگا (لەگەڵ بۆتەکە بار دەکرێت — هەمیشە بەردەستە)،
# پاشان ڕێڕەوی سیستەم (لە Actions بە apt دامەزرێت)، لە کۆتاییدا داگرتن لە ئینتەرنێت
REPO_FONTS_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "fonts")
FONT_CANDIDATES = [
    os.path.join(REPO_FONTS_DIR, "NotoNaskhArabic-Regular.ttf"),
    os.path.join(REPO_FONTS_DIR, "Amiri-Regular.ttf"),
    "/usr/share/fonts/truetype/noto/NotoNaskhArabic-Regular.ttf",
    "/usr/share/fonts/opentype/noto/NotoNaskhArabic-Regular.ttf",
    "/usr/share/fonts/truetype/amiri/Amiri-Regular.ttf",
    "/usr/share/fonts/truetype/noto/NotoSansArabic-Regular.ttf",
]
FONT_URLS = [
    "https://github.com/google/fonts/raw/main/ofl/amiri/Amiri-Regular.ttf",
    "https://raw.githubusercontent.com/notofonts/notofonts.github.io/main/fonts/NotoNaskhArabic/hinted/ttf/NotoNaskhArabic-Regular.ttf",
    "https://cdn.jsdelivr.net/gh/notofonts/notofonts.github.io/fonts/NotoNaskhArabic/hinted/ttf/NotoNaskhArabic-Regular.ttf",
]


def ensure_arabic_font(workdir):
    for p in FONT_CANDIDATES:
        if os.path.exists(p):
            return p
    fdir = os.path.join(workdir, "fonts")
    os.makedirs(fdir, exist_ok=True)
    fp = os.path.join(fdir, "arabic.ttf")
    if os.path.exists(fp) and os.path.getsize(fp) > 100_000:
        return fp            # فۆنتەکە پێشتر دابەزیوە — دووبارە داگری ناکەین
    for url in FONT_URLS:
        try:
            r = SESSION.get(url, timeout=90)
            r.raise_for_status()
            if len(r.content) > 100_000:
                with open(fp, "wb") as fh:
                    fh.write(r.content)
                log.info("downloaded Arabic font from %s", url)
                return fp
        except Exception as e:
            log.warning("font download failed %s: %s", url, e)
    return None


def extract_pdf_text(data):
    doc = fitz_mod.open(stream=data, filetype="pdf")
    pages = []
    for pg in doc:
        try:
            pages.append(pg.get_text("text") or "")
        except Exception:
            pages.append("")
    doc.close()
    if not has_letters("".join(pages)):
        raise BotError(ERR_SCANNED)
    return pages


def ocr_pdf_pages(data, workdir, dpi=170, max_pages=40):
    """OCR بۆ PDF ی سکانیکراو بە tesseract (ara+eng) — بۆ ئەو PDF اIANE کە دەقیان نییە.
    None ئەگەر tesseract بەردەست نییە یان دەقی گونجاو دەرنەچوو."""
    if not shutil.which("tesseract"):
        log.info("tesseract not available — OCR skipped")
        return None
    doc = fitz_mod.open(stream=data, filetype="pdf")
    tmpdir = tempfile.mkdtemp(prefix="ocr_", dir=workdir)
    pages = []
    try:
        for i, pg in enumerate(doc):
            if i >= max_pages:
                break
            png = os.path.join(tmpdir, "pg%03d.png" % i)
            try:
                pg.get_pixmap(dpi=dpi).save(png)
                res = subprocess.run(["tesseract", png, "stdout", "-l", "ara+eng"],
                                     capture_output=True, text=True, timeout=240)
                pages.append((res.stdout or "") if res.returncode == 0 else "")
            except Exception:
                pages.append("")
            finally:
                try:
                    os.remove(png)
                except Exception:
                    pass
    finally:
        doc.close()
        shutil.rmtree(tmpdir, ignore_errors=True)
    if not has_letters("".join(pages)):
        return None
    log.info("OCR extracted %d chars from %d pages",
             sum(len(p) for p in pages), len(pages))
    return pages


def extract_pdf_text_or_ocr(data, workdir):
    """دەرهێنانی دەقی PDF؛ ئەگەر سکانیکراو بوو یان دەق زۆر کەم بوو، خۆکارانە OCR تاقی دەکاتەوە."""
    pages = None
    try:
        pages = extract_pdf_text(data)
    except BotError:
        pages = None
    if pages:
        avg = sum(len(p.strip()) for p in pages) / max(1, len(pages))
        if avg < 40:      # دەق زۆر کەمە بەرامبەر پەڕەکان — لەوانەیە سکانیکراوە
            try:
                ocr = ocr_pdf_pages(data, workdir)
                if ocr and sum(len(x) for x in ocr) > sum(len(x) for x in pages):
                    pages = ocr
            except Exception as e:
                log.warning("OCR fallback failed: %s", e)
        return pages
    try:
        pages = ocr_pdf_pages(data, workdir)
    except Exception as e:
        log.warning("OCR failed: %s", e)
    if not pages:
        raise BotError(ERR_SCANNED)
    return pages


def _wrap_rtl_lines(text, font_name, size, maxw):
    """دابەشکردنی دەق بۆ دێری هاوسەنگ، پاشان شێوەدان + bidi بۆ هەر دێرێک."""
    lines = []
    for raw in text.splitlines():
        raw = raw.strip()
        if not raw:
            lines.append("")
            continue
        words = raw.split()
        cur = ""
        for w in words:
            cand = (cur + " " + w) if cur else w
            try:
                width = pdfmetrics.stringWidth(_shape(cand), font_name, size)
            except Exception:
                width = len(cand) * size * 0.55
            if width <= maxw:
                cur = cand
            elif not cur:
                # وشەیەکی زۆر درێژ (لینک...) — بە زۆر دەبڕین
                chunk = w
                while chunk:
                    k = len(chunk)
                    while k > 1 and pdfmetrics.stringWidth(_shape(chunk[:k]), font_name, size) > maxw:
                        k -= 1
                    lines.append(_disp(_shape(chunk[:k])))
                    chunk = chunk[k:].lstrip()
                cur = ""
            else:
                lines.append(_disp(_shape(cur)))
                cur = w
        if cur:
            lines.append(_disp(_shape(cur)))
    return lines


def build_pdf(pages_texts, out_path, title="Sorani Translation"):
    workdir = os.path.dirname(os.path.abspath(out_path)) or "."
    font_path = ensure_arabic_font(workdir)
    if not font_path:
        raise RuntimeError("no Arabic font available (download failed)")
    pdfmetrics.registerFont(TTFont("KurdishFont", font_path))
    W, H = A4
    M = 46.0
    body, lead = 12.5, 21.0
    maxw = W - 2 * M
    c = rl_canvas.Canvas(out_path, pagesize=A4)
    c.setTitle((title or "Sorani Translation")[:120])
    c.setAuthor("Sorani Translator Bot")
    for i, ptxt in enumerate(pages_texts, 1):
        c.setFont("KurdishFont", 9)
        c.setFillColorRGB(0.45, 0.45, 0.45)
        c.drawRightString(W - M, H - M + 8, _disp(_shape("پەڕەی %d" % i)))
        c.setFillColorRGB(0, 0, 0)
        y = H - M - 26
        for ln in _wrap_rtl_lines(ptxt, "KurdishFont", body, maxw):
            if y < M + lead:
                c.showPage()
                c.setFont("KurdishFont", body)
                c.setFillColorRGB(0, 0, 0)
                y = H - M
            if ln:
                c.setFont("KurdishFont", body)
                c.drawRightString(W - M, y, ln)
            y -= lead
        c.showPage()
    c.save()
    return out_path


# ---------- PDF ی گشتی: بلۆکەکان + گۆڕین بە LibreOffice (کۆپی PDF) ----------

_PDF_STYLES = {
    "title":  (16.0, 24.0, 10.0),
    "h1":     (15.0, 22.0, 8.0),
    "h2":     (13.5, 20.0, 6.0),
    "body":   (12.5, 21.0, 4.0),
    "bullet": (12.5, 21.0, 3.0),
    "meta":   (10.5, 17.0, 3.0),
}


def build_pdf_blocks(blocks, out_path, title="Sorani Translation"):
    """PDF ی RTL لە لیستی بلۆک (جۆر، دەق). جۆرەکان: title/h1/h2/body/bullet/meta"""
    workdir = os.path.dirname(os.path.abspath(out_path)) or "."
    font_path = ensure_arabic_font(workdir)
    if not font_path:
        raise RuntimeError("no Arabic font available (download failed)")
    pdfmetrics.registerFont(TTFont("KurdishFont", font_path))
    W, H = A4
    M = 46.0
    maxw = W - 2 * M
    c = rl_canvas.Canvas(out_path, pagesize=A4)
    c.setTitle((title or "Sorani Translation")[:120])
    c.setAuthor("Sorani Translator Bot")

    pg = {"n": 1}

    def draw_footer():
        c.setFont("KurdishFont", 9)
        c.setFillColorRGB(0.45, 0.45, 0.45)
        c.drawRightString(W - M, H - M + 8, _disp(_shape("پەڕەی %d" % pg["n"])))
        c.setFillColorRGB(0, 0, 0)

    def start_page():
        c.showPage()
        pg["n"] += 1
        draw_footer()
        return H - M

    draw_footer()
    y = H - M - 10
    if title:
        for ln in _wrap_rtl_lines(title, "KurdishFont", 16.0, maxw):
            c.setFont("KurdishFont", 16.0)
            c.drawRightString(W - M, y, ln)
            y -= 24.0
        y -= 10.0

    for kind, text in blocks:
        text = (text or "").strip()
        if not text:
            continue
        size, lead, gap = _PDF_STYLES.get(kind, _PDF_STYLES["body"])
        if kind == "bullet":
            text = "• " + text
        for ln in _wrap_rtl_lines(text, "KurdishFont", size, maxw):
            if y < M + lead:
                y = start_page()
            if ln:
                c.setFont("KurdishFont", size)
                c.setFillColorRGB(0.1, 0.1, 0.1) if kind in ("title", "h1", "h2") \
                    else c.setFillColorRGB(0, 0, 0)
                c.drawRightString(W - M, y, ln)
            y -= lead
        y -= gap
    c.save()
    return out_path


def _soffice_bin():
    return shutil.which("soffice") or shutil.which("libreoffice")


def office_to_pdf(src_path, out_dir, timeout=420):
    """گۆڕینی DOCX/PPTX بۆ PDF بە LibreOffice — هەمان شێواز و وێنەکان دەمێننەوە.
    ئەگەر نەکرا None دەگەڕێنێتەوە (کۆدەکە بە ReportLab بەردەوام دەبێت)."""
    soff = _soffice_bin()
    if not soff:
        log.info("soffice not available — using ReportLab fallback")
        return None
    out_dir = os.path.abspath(out_dir)
    os.makedirs(out_dir, exist_ok=True)
    prof = tempfile.mkdtemp(prefix="lo_profile_")
    cmd = [soff, "-env:UserInstallation=file://" + prof,
           "--headless", "--norestore", "--nolockcheck",
           "--convert-to", "pdf", "--outdir", out_dir, os.path.abspath(src_path)]
    try:
        res = subprocess.run(cmd, capture_output=True, text=True, timeout=timeout)
    except Exception as e:
        log.warning("soffice failed: %s", e)
        return None
    finally:
        shutil.rmtree(prof, ignore_errors=True)
    stem = os.path.splitext(os.path.basename(src_path))[0]
    out = os.path.join(out_dir, stem + ".pdf")
    if res.returncode != 0 or not os.path.exists(out) or os.path.getsize(out) == 0:
        log.warning("soffice convert failed: rc=%s %s",
                    res.returncode, (res.stderr or "")[-300:])
        return None
    log.info("LibreOffice converted %s -> %s", os.path.basename(src_path), os.path.basename(out))
    return out


def register_font_fontconfig(font_path):
    """فۆنتەکە دەخاتە ~/.fonts بۆ ئەوەی LibreOffice بیبینێت (بۆ نووسینی عەرەبی/کوردی)."""
    try:
        fdir = os.path.expanduser("~/.fonts")
        os.makedirs(fdir, exist_ok=True)
        dst = os.path.join(fdir, "bot-arabic.ttf")
        if os.path.abspath(font_path) != os.path.abspath(dst):
            shutil.copyfile(font_path, dst)
        subprocess.run(["fc-cache", "-f"], capture_output=True, timeout=90)
    except Exception as e:
        log.warning("fontconfig register failed: %s", e)


def ensure_font_and_register(workdir):
    fp = ensure_arabic_font(workdir)
    if fp:
        register_font_fontconfig(fp)
    return fp


def _docx_para_block(p):
    t = (p.text or "").strip()
    if not t:
        return []
    try:
        sname = ((p.style and p.style.name) or "").lower()
    except Exception:
        sname = ""
    if sname == "title" or "heading 1" in sname:
        return [("h1", t)]
    if "heading" in sname:
        return [("h2", t)]
    if "list" in sname or "bullet" in sname:
        return [("bullet", t)]
    return [("body", t)]


def blocks_from_docx(path):
    """بلۆکەکان لە فایلی Word (پاراگراف + خشتە بە هەمان ڕیزبەندی)."""
    doc = DocxDocument(path)
    blocks = []
    try:
        from docx.oxml.table import CT_Tbl
        from docx.oxml.text.paragraph import CT_P
        from docx.table import Table as _Tbl
        from docx.text.paragraph import Paragraph as _Para
        for child in doc.element.body.iterchildren():
            if isinstance(child, CT_P):
                blocks.extend(_docx_para_block(_Para(child, doc)))
            elif isinstance(child, CT_Tbl):
                try:
                    for row in _Tbl(child, doc).rows:
                        cells = []
                        for cell in row.cells:
                            ct = " ".join(x.strip() for x in (cell.text or "").splitlines()
                                          if x.strip())
                            if ct:
                                cells.append(ct)
                        if cells:
                            blocks.append(("body", " | ".join(cells)))
                except Exception:
                    pass
    except Exception:
        for p in doc.paragraphs:
            blocks.extend(_docx_para_block(p))
    return [b for b in blocks if (b[1] or "").strip()]


def blocks_from_pptx(path):
    """بلۆکەکان لە فایلی PowerPoint (سلاید بە سلاید)."""
    prs = Presentation(path)
    blocks = []
    for i, slide in enumerate(prs.slides, 1):
        slide_blocks = []
        for sh in _iter_shapes(slide.shapes):
            try:
                if sh.has_text_frame:
                    for para in _ppt_paras(sh.text_frame):
                        t = (para.text or "").strip()
                        if t and has_letters(t):
                            slide_blocks.append(
                                ("bullet" if getattr(para, "level", 0) else "body", t))
            except Exception:
                pass
        if slide_blocks:
            blocks.append(("h2", "سلایدی %d" % i))
            blocks.extend(slide_blocks)
    return blocks


def blocks_from_plain(text):
    blocks = []
    for ln in (text or "").splitlines():
        ln = ln.strip()
        if ln:
            blocks.append(("body", ln))
    return blocks


def blocks_from_transcript(cues, texts):
    """بلۆکی نووسینی دەنگ: cues دەکرێت (stamp_str, text) یان (st, en, text) بێت."""
    blocks = []
    for cue, t in zip(cues, texts):
        t = (t or "").strip()
        if not t:
            continue
        raw = cue[0] if len(cue) == 2 else ts(cue[0])
        stamp = str(raw).split("-->")[0].strip()
        blocks.append(("meta", "[%s]" % stamp))
        blocks.append(("body", t))
    return blocks


def docx_to_pdf(docx_path, workdir, title):
    """PDF بۆ Word: سەرەتا LibreOffice (شێوازی ڕەسەن)، ئەگەر نەبوو ReportLab."""
    out = None
    try:
        out = office_to_pdf(docx_path, workdir)
    except Exception as e:
        log.warning("office_to_pdf(docx) error: %s", e)
    if out:
        return out
    try:
        blocks = blocks_from_docx(docx_path)
        if not blocks:
            return None
        stem = sanitize(os.path.splitext(os.path.basename(docx_path))[0])
        return build_pdf_blocks(blocks, os.path.join(workdir, stem + ".pdf"), title)
    except Exception as e:
        log.warning("docx pdf fallback failed: %s", e)
        return None


def pptx_to_pdf(pptx_path, workdir, title):
    """PDF بۆ PowerPoint: سەرەتا LibreOffice، ئەگەر نەبوو ReportLab."""
    out = None
    try:
        out = office_to_pdf(pptx_path, workdir)
    except Exception as e:
        log.warning("office_to_pdf(pptx) error: %s", e)
    if out:
        return out
    try:
        blocks = blocks_from_pptx(pptx_path)
        if not blocks:
            return None
        stem = sanitize(os.path.splitext(os.path.basename(pptx_path))[0])
        return build_pdf_blocks(blocks, os.path.join(workdir, stem + ".pdf"), title)
    except Exception as e:
        log.warning("pptx pdf fallback failed: %s", e)
        return None


# ------------------- دەنگی کوردی سۆرانی (TTS) -----------------------
#  kurdishtts.com — POST /api/tts-proxy بە x-api-key → دەنگی سۆرانی (wav/mp3)
#  پلانی خۆڕایی: 20,000 پیت/مانگ، 500 پیت بۆ هەر داواکاری — دەنگی خۆڕایی: sorani_85

_TTS_MEMO = {}


def tts_available():
    return bool(TTS_API_KEY)


def _tts_once(text):
    if not TTS_API_KEY:
        raise RuntimeError("KurdishTTS not configured (set KURDISH_TTS_API_KEY secret)")
    r = SESSION.post(
        TTS_API_URL,
        headers={"x-api-key": TTS_API_KEY, "Content-Type": "application/json"},
        json={"text": text[:490], "speaker_id": TTS_SPEAKER,
              "model_version": TTS_MODEL_V, "format": "mp3",
              "speed": TTS_SPEED, "dialect": "sorani"},
        timeout=120,
    )
    if r.status_code in (401, 403):
        raise RuntimeError("kurdishtts http %s — key/quota: %s"
                           % (r.status_code, (r.text or "")[:120]))
    if r.status_code != 200:
        raise RuntimeError("kurdishtts http %s: %s" % (r.status_code, (r.text or "")[:120]))
    if len(r.content) < 512:
        raise RuntimeError("kurdishtts tiny response (%d bytes)" % len(r.content))
    return r.content


def tts_clip(text):
    """دەنگی پارچەیەک بە یادەوەری — None ئەگەر سەرنەچوو."""
    key = (text or "").strip()
    if not key:
        return None
    if key in _TTS_MEMO:
        return _TTS_MEMO[key]
    try:
        data = _tts_once(key)
    except Exception as e:
        log.warning("tts failed: %s", e)
        _TTS_MEMO[key] = None
        return None
    _TTS_MEMO[key] = data
    return data


def _atempo_chain(ratio):
    """زنجیرەی atempo بۆ خێراکردنی دەنگ بۆ ratio — هەر پارچەیەک ≤ 2.0،
    کۆی گشتی سنووردارە بە ٢.٥x بۆ ئەوەی دەنگ هێشتا سروشتی بمێنێتەوە."""
    ratio = min(max(ratio, 0.5), 2.5)
    parts, r = [], ratio
    while r > 2.0 + 1e-6:
        parts.append("atempo=2.0")
        r /= 2.0
    parts.append("atempo=%.3f" % r)
    return ",".join(parts)


def _fit_voice_clip(src, gap, dst):
    """کلیپی TTS دەگونجێنێت لە مەودای بەردەستدا (لە دەستپێکی خۆی تا دەستپێکی
    پارچەی داهاتوو) — بۆ هاوکاتی دروست لەگەڵ ڤیدیۆ:
    - ئەگەر کلیپەکە درێژتر بوو لە مەوداکە، بە atempo خێرا دەکرێت (تا ٢.٥x)
    - ئەگەر هێشتا درێژ بوو، لە کۆتایی مەوداکە دەبڕدرێت بە fadeout
    دەنگ ٤٤.١کHZ ستێریۆ دەگەڕێنرێتەوە بۆ یەکخستنی پاک."""
    d = ffprobe_duration(src)
    filters = []
    if gap and gap >= 0.6 and d > gap:
        eff = gap - 0.15                       # پاشگەزبەندی بچووک پێش پارچەی داهاتوو
        ratio = d / max(eff, 0.5)
        filters.append(_atempo_chain(ratio))
        nd = d / min(ratio, 2.5)
        if nd > eff + 0.02:                    # هێشتا درێژە — دەبڕین بە fadeout
            filters.append("atrim=0:%.3f" % eff)
            f0 = max(0.0, eff - 0.30)
            filters.append("afade=t=out:st=%.3f:d=%.3f" % (f0, eff - f0))
    elif gap and 0 < gap < 0.6 and d > gap:
        filters.append("atrim=0:%.3f" % max(0.3, gap - 0.05))
        filters.append("afade=t=out:st=0:d=0.15")
    cmd = ["ffmpeg", "-y", "-loglevel", "error", "-i", src,
           "-af", ",".join(filters) if filters else "anull",
           "-ac", "2", "-ar", "44100", dst]
    try:
        res = subprocess.run(cmd, capture_output=True, text=True, timeout=300)
    except Exception as e:
        log.warning("fit voice clip error: %s", e)
        return False
    if res.returncode != 0 or not os.path.exists(dst) or os.path.getsize(dst) == 0:
        log.warning("fit voice clip failed: %s", (res.stderr or "")[-300:])
        return False
    return True


def build_kurdish_voice_track(cues, texts, total_dur, workdir, workers=3):
    """تراکێکی دەنگی سۆرانی لە ژێرنووسە وەرگێڕدراوەکان دروست دەکات.

    **هاوکاتی تەواو لەگەڵ ڤیدیۆ:** هەر پارچەیەک ڕێک لە چرکەی ژێرنووسەکەی خۆی
    دەست پێدەکات (adelay بە پانتایی چرکە) — واتە ئەگەر لە چرکەی ٢٠ وشەیەک لە
    ڤیدیۆکە بوترێت، وەرگێڕانە کوردییەکەشی لە هەمان چرکە دەست پێدەکات.
    ئەگەر پارچەیەک درێژتر بوو لە مەودای بەردەست (تا دەستپێکی پارچەی داهاتوو)،
    خۆکارانە خێراتر دەکرێت (تا ٢.٥x) یان دەبڕدرێت — هەرگیز لەسەر یەک ناکەون.
    ئەگەر هیچ کلیپێک نەبوو None."""
    if not shutil.which("ffmpeg"):
        raise BotError(ERR_FFMPEG)
    tdir = tempfile.mkdtemp(prefix="tts_", dir=workdir)
    items, used = [], 0
    for i, ((st, en, _), t) in enumerate(zip(cues, texts)):
        t = (t or "").strip()
        if not t:
            continue
        if used + len(t) > TTS_MAX_CHARS:
            log.info("tts char budget reached (%d/%d) — remaining cues skipped",
                     used, TTS_MAX_CHARS)
            break
        used += len(t)
        items.append((i, float(st), t))
    if not items:
        return None

    def _get(item):
        i, st, t = item
        b = tts_clip(t)
        if not b:
            return None
        p = os.path.join(tdir, "c%04d.mp3" % i)
        try:
            with open(p, "wb") as fh:
                fh.write(b)
        except Exception:
            return None
        return (st, p)

    clips = []
    with ThreadPoolExecutor(max(1, min(workers, len(items)))) as ex:
        for r in ex.map(_get, items):
            if r:
                clips.append(r)
    if not clips:
        return None
    clips.sort(key=lambda x: x[0])

    total = max(float(total_dur or 0), clips[-1][0] + 2.0) + 1.0
    # ---- گونجاندنی هەر کلیپێک لە مەودای خۆیدا (sync + بێ لەسەریەککەوتن) ----
    fitted = []
    for n, (st, p) in enumerate(clips):
        if n + 1 < len(clips):
            gap = max(0.3, clips[n + 1][0] - st)
        else:
            gap = max(2.0, total - 0.5 - st)
        fw = os.path.join(tdir, "f%04d.wav" % n)
        if _fit_voice_clip(p, gap, fw):
            fitted.append((st, fw))
        else:
            fitted.append((st, p))          # پاشەکەوت: فایلی ڕەسەن بەکاردەهێنرێت
    if not fitted:
        return None

    sil = os.path.join(tdir, "sil.wav")
    try:
        res = subprocess.run(
            ["ffmpeg", "-y", "-loglevel", "error", "-f", "lavfi",
             "-i", "anullsrc=r=44100:cl=stereo", "-t", "%.2f" % total, sil],
            capture_output=True, text=True, timeout=300)
        if res.returncode != 0:
            return None
        inputs = ["-i", sil]
        parts, mix = [], "[0:a]"
        for idx, (st, p) in enumerate(fitted, 1):
            inputs += ["-i", p]
            d = int(st * 1000)
            parts.append("[%d:a]aresample=44100,adelay=%d:all=1[a%d]"
                         % (idx, d, idx))
            mix += "[a%d]" % idx
        parts.append("%samix=inputs=%d:duration=first:normalize=0[out]"
                     % (mix, len(fitted) + 1))
        out = os.path.join(tdir, "voice.wav")
        res = subprocess.run(
            ["ffmpeg", "-y", "-loglevel", "error"] + inputs +
            ["-filter_complex", ";".join(parts), "-map", "[out]",
             "-c:a", "pcm_s16le", out],
            capture_output=True, text=True, timeout=1800)
        if res.returncode != 0 or not os.path.exists(out) or os.path.getsize(out) == 0:
            log.warning("voice track failed: %s", (res.stderr or "")[-400:])
            return None
        log.info("voice track built: %d clips, %.1fs (synced to cue starts)",
                 len(fitted), total)
        return out
    except Exception as e:
        log.warning("voice track error: %s", e)
        return None


# ----------------------- ڤیدیۆ / دەنگ -----------------------------

def ffmpeg_extract(src, dst):
    if not shutil.which("ffmpeg"):
        raise BotError(ERR_FFMPEG)
    cmd = ["ffmpeg", "-y", "-loglevel", "error", "-i", src,
           "-vn", "-ac", "1", "-ar", "16000", dst]
    try:
        res = subprocess.run(cmd, capture_output=True, text=True, timeout=1200)
    except subprocess.TimeoutExpired:
        raise BotError(ERR_LONG)
    if res.returncode != 0 or not os.path.exists(dst):
        log.warning("ffmpeg failed: %s", (res.stderr or "")[-500:])
        raise BotError(ERR_AUDIO)


def ffprobe_duration(path):
    if not shutil.which("ffprobe"):
        return 0.0
    try:
        res = subprocess.run(
            ["ffprobe", "-v", "error", "-show_entries", "format=duration",
             "-of", "default=nw=1:nk=1", path],
            capture_output=True, text=True, timeout=60)
        return float(res.stdout.strip()) if res.returncode == 0 else 0.0
    except Exception:
        return 0.0


def ffprobe_has_video(path):
    """ئایا فایلەکە تراکێکی ڤیدیۆی تێدایە؟ (بۆ دەنگی خاڵیس False)"""
    if not shutil.which("ffprobe"):
        return False
    try:
        res = subprocess.run(
            ["ffprobe", "-v", "error", "-select_streams", "v:0",
             "-show_entries", "stream=codec_type", "-of", "csv=p=0", path],
            capture_output=True, text=True, timeout=60)
        return "video" in (res.stdout or "")
    except Exception:
        return False


def ffprobe_dims(path):
    """(پان, بەرزی) ی تراکی ڤیدیۆ — ئەگەر نەبوو (0, 0)."""
    if not shutil.which("ffprobe"):
        return 0, 0
    try:
        res = subprocess.run(
            ["ffprobe", "-v", "error", "-select_streams", "v:0",
             "-show_entries", "stream=width,height", "-of", "csv=p=0:s=x", path],
            capture_output=True, text=True, timeout=60)
        for line in (res.stdout or "").splitlines():
            line = line.strip()
            if "x" in line:
                a, b = line.split("x", 1)
                try:
                    return int(float(a)), int(float(b))
                except ValueError:
                    continue
        return 0, 0
    except Exception:
        return 0, 0


def compress_video(src, dst, target_mb, workdir, timeout=2400):
    """بچووککردنەوەی ڤیدیۆ بۆ خوار target_mb مێگابایت — دوو-پاس بە bitrate ی ژمێریارکراو.
    ئەگەر وێنەکە گەورەیە و bitrate کەمە، خۆکارانە بۆ 720p/480p دادەبەزێت.
    ئەگەر سەرکەوتوو نەبوو None دەگەڕێنێتەوە."""
    dur = ffprobe_duration(src)
    if dur <= 0:
        return None
    target_mb = max(0.3, float(target_mb))
    audio_kbps = 96
    v_kbps = int(target_mb * 1024 * 1024 * 8.0 / dur / 1000.0 * 0.96) - audio_kbps
    v_kbps = max(100, v_kbps)
    _, h = ffprobe_dims(src)
    vf = None
    if h:
        if v_kbps < 350 and h > 480:
            vf = "scale=-2:480"
        elif v_kbps < 900 and h > 720:
            vf = "scale=-2:720"
    passlog = os.path.join(workdir, "x264pass")

    def _two_pass(vk, out):
        base = ["ffmpeg", "-y", "-loglevel", "error", "-i", src]
        if vf:
            base += ["-vf", vf]
        opts = ["-c:v", "libx264", "-preset", "veryfast",
                "-b:v", "%dk" % vk,
                "-maxrate", "%dk" % int(vk * 1.45),
                "-bufsize", "%dk" % (vk * 3),
                "-c:a", "aac", "-b:a", "%dk" % audio_kbps]
        cmd1 = base + opts + ["-pass", "1", "-passlogfile", passlog,
                              "-f", "null", "/dev/null"]
        cmd2 = base + opts + ["-pass", "2", "-passlogfile", passlog,
                              "-movflags", "+faststart", out]
        for c in (cmd1, cmd2):
            try:
                res = subprocess.run(c, capture_output=True, text=True, timeout=timeout)
            except subprocess.TimeoutExpired:
                return False
            if res.returncode != 0:
                log.warning("compress pass failed: %s", (res.stderr or "")[-300:])
                return False
        return os.path.exists(out) and os.path.getsize(out) > 0

    try:
        ok = _two_pass(v_kbps, dst)
        if ok and os.path.getsize(dst) > target_mb * 1024 * 1024 * 1.02:
            ok = _two_pass(int(v_kbps * 0.8), dst)   # هەوڵێکی تر بە bitrate ی نزمتر
    finally:
        for f in glob.glob(passlog + "*"):
            try:
                os.remove(f)
            except Exception:
                pass
    return dst if ok else None


def shrink_source_for_processing(path, workdir):
    """ئەگەر ڤیدیۆ زۆر گەورەیە (لە SHRINK_OVER_MB زیاتر یان بەرزی > 720)،
    پێش وەرگێڕان و سووتاندنی ژێرنووس بۆ 720p بچووکی دەکرێتەوە — خێراتر کار دەکات.
    ئەگەر بچووککردنەوە سوود نەبوو، فایلی ڕەسەن دەگەڕێنێتەوە."""
    try:
        if not shutil.which("ffmpeg"):
            return path
        size = os.path.getsize(path)
        _, h = ffprobe_dims(path)
        if size <= SHRINK_OVER_MB * 1024 * 1024 and (not h or h <= 720):
            return path
        out = os.path.join(os.path.dirname(path) or workdir, "src_small.mp4")
        cmd = ["ffmpeg", "-y", "-loglevel", "error", "-i", path]
        if h and h > 720:
            cmd += ["-vf", "scale=-2:720"]
        cmd += ["-c:v", "libx264", "-crf", "26", "-preset", "veryfast",
                "-c:a", "aac", "-b:a", "128k", out]
        res = subprocess.run(cmd, capture_output=True, text=True, timeout=3000)
        if (res.returncode == 0 and os.path.exists(out) and os.path.getsize(out) > 0
                and os.path.getsize(out) < size):
            log.info("source shrunk: %.1fMB -> %.1fMB",
                     size / 1048576.0, os.path.getsize(out) / 1048576.0)
            try:
                os.remove(path)
            except Exception:
                pass
            return out
        log.warning("shrink did not help — keeping original")
    except Exception as e:
        log.warning("shrink failed: %s", e)
    return path


_FONT_FAMILY_MAP = {
    # ناوی فایل -> ناوی ڕاستەقینەی خێزان (ئەگەر fontTools بەردەست نەبوو)
    "notonaskharabic": "Noto Naskh Arabic",
    "notosansarabic": "Noto Sans Arabic",
    "notosansarabicui": "Noto Sans Arabic UI",
    "amiri": "Amiri",
    "arabic": "Amiri",     # فۆنتی داگیراو بە ناوی arabic.ttf — ناوەڕاستەکەی Amiri ـە
}


def _ass_font_family(font_path):
    """ناوی ڕاستەقینەی خێزانی فۆنت لە ناو فایلەکەوە دەخوێنێتەوە (name table).
    گرنگ نییە فایلەکە بە چ ناوێک هەڵگیراوە — ناکۆکی ناوی FontName بووە
    هۆی ئەوەی libass فۆنتی هەڵە بەکاربهێنێت و پیتەکان جیابنەوە."""
    try:
        from fontTools.ttLib import TTFont
        tt = TTFont(font_path, fontNumber=0, lazy=True)
        try:
            names = tt["name"]
            for nid in (16, 1):     # 16 = Typographic Family, 1 = Family
                v = names.getDebugName(nid)
                if v and v.strip():
                    return v.strip()
        finally:
            tt.close()
    except Exception as e:
        log.debug("fontTools family read failed: %s", e)
    stem = os.path.splitext(os.path.basename(font_path or ""))[0]
    stem = stem.split("-")[0]
    key = re.sub(r"[^a-z0-9]", "", stem.lower())
    if key in _FONT_FAMILY_MAP:
        return _FONT_FAMILY_MAP[key]
    fam = re.sub(r"(?<=[a-z0-9])(?=[A-Z])", " ", stem)
    return fam.strip() or "Amiri"


def _ass_style(fam):
    return ("FontName=%s,Fontsize=22,PrimaryColour=&H00FFFFFF,"
            "OutlineColour=&H00000000,BackColour=&H80000000,BorderStyle=1,"
            "Outline=2,Shadow=1,Alignment=2,MarginV=28,Spacing=0" % fam)


def _ff_quote(s):
    """ڕێڕەو بۆ ناو filtergraph ی ffmpeg — بە ' ' دەپەچێت (پارێزراو لە , و :)"""
    return "'" + s.replace("\\", "\\\\").replace("'", "\\'") + "'"


def _srt_ts_to_ass(s):
    """00:00:01,240 -> 0:00:01.24  (فۆرماتی ASS)"""
    hh, mm, rest = s.strip().split(":")
    rest = rest.replace(",", ".")
    if "." in rest:
        ss, ms = rest.split(".", 1)
    else:
        ss, ms = rest, "0"
    return "%d:%s:%s.%s" % (int(hh), mm, ss, (ms + "00")[:2])


def build_ass_from_srt(srt_path, ass_path, family):
    """SRT دەکاتە فایلی ASS ی دروست بە ستایلی ڕوون و فۆنتی ڕاستەقینە.
    libass بە HarfBuzz شێوەی پیتەکان دروست دەکات و بە FriBidi ڕاستەوخۆ RTL دەکات،
    بۆیە دەقەکە بە هەمان باری لۆجیکی خۆی دەنووسرێت (بێ reshaper/bidi)."""
    with open(srt_path, "r", encoding="utf-8") as fh:
        raw = fh.read()
    events = []
    for block in re.split(r"\n\s*\n", raw.strip()):
        lines = [x.strip() for x in block.splitlines() if x.strip()]
        if len(lines) < 2:
            continue
        t_idx = None
        for i, ln in enumerate(lines):
            if re.match(r"\d{1,2}:\d{2}:\d{2}[,.]\d{1,3}\s*-->\s*\d{1,2}:\d{2}:\d{2}[,.]\d{1,3}", ln):
                t_idx = i
                break
        if t_idx is None:
            continue
        m = re.match(r"(\d{1,2}:\d{2}:\d{2}[,.]\d{1,3})\s*-->\s*(\d{1,2}:\d{2}:\d{2}[,.]\d{1,3})",
                     lines[t_idx])
        txt = "\\N".join(re.sub(r"[{}\\]", "", x) for x in lines[t_idx + 1:]).strip()
        if not txt:
            continue
        events.append("Dialogue: 0,%s,%s,Kur,,0,0,0,,%s"
                      % (_srt_ts_to_ass(m.group(1)), _srt_ts_to_ass(m.group(2)), txt))
    if not events:
        return None
    header = (
        "[Script Info]\n"
        "; Sorani subtitles — burn-in (generated by bot.py)\n"
        "ScriptType: v4.00+\n"
        "PlayResX: 384\n"
        "PlayResY: 288\n"
        "WrapStyle: 0\n"
        "ScaledBorderAndShadow: yes\n"
        "YCbCr Matrix: TV.709\n\n"
        "[V4+ Styles]\n"
        "Format: Name, Fontname, Fontsize, PrimaryColour, SecondaryColour, "
        "OutlineColour, BackColour, Bold, Italic, Underline, StrikeOut, ScaleX, ScaleY, "
        "Spacing, Angle, BorderStyle, Outline, Shadow, Alignment, MarginL, MarginR, "
        "MarginV, Encoding\n"
        "Style: Kur,%s,22,&H00FFFFFF,&H000000FF,&H00000000,&H7F000000,"
        "0,0,0,0,100,100,0,0,1,2,1,2,12,12,28,1\n\n"
        "[Events]\n"
        "Format: Layer, Start, End, Style, Name, MarginL, MarginR, MarginV, Effect, Text\n"
        % family)
    with open(ass_path, "w", encoding="utf-8") as fh:
        fh.write(header + "\n".join(events) + "\n")
    return ass_path


def burn_subtitles(src_video, srt_path, out_path, workdir, timeout=2400, voice_track=None):
    """ژێرنووسی کوردی دەکاتەوە سەر خۆی ڤیدیۆکە (ffmpeg + libass).
    ئەگەر voice_track درابێت، دەنگی کوردی جیگای دەنگی ڕەسەن دادەنرێت.
    فۆنتی عەرەبی/کوردی لە fontsdir دا دەدۆزرێتەوە. ئەگەر سەرنەچوو None."""
    if not shutil.which("ffmpeg"):
        raise BotError(ERR_FFMPEG)
    font_path = ensure_arabic_font(workdir)
    cwd = os.path.abspath(workdir)
    srt_local = os.path.join(cwd, "burn_subs.srt")
    shutil.copyfile(srt_path, srt_local)
    ext = os.path.splitext(src_video)[1] or ".mp4"
    src_local = os.path.join(cwd, "burn_src" + ext)
    if os.path.abspath(src_video) != os.path.abspath(src_local):
        shutil.copyfile(src_video, src_local)
    fontsdir = (os.path.dirname(os.path.abspath(font_path))
                if font_path else "/usr/share/fonts")
    # ١) ڕێگای سەرەکی: فایلی ASS بە ناوی ڕاستەقینەی فۆنت — شێوەدانی پیتەکان
    #    بە HarfBuzz و RTL بە FriBidi دەکرێت (پیتەکان بە دروستی بەستراوە دەردەکەون)
    vf = None
    if font_path:
        fam = _ass_font_family(font_path)
        ass_local = os.path.join(cwd, "burn_subs.ass")
        try:
            if build_ass_from_srt(srt_local, ass_local, fam):
                vf = "subtitles=burn_subs.ass:fontsdir=%s" % _ff_quote(fontsdir)
                log.info("burn: ASS subtitles with font family '%s' @ %s", fam, font_path)
        except Exception as e:
            log.warning("ASS build failed, falling back to SRT: %s", e)
    if vf is None:
        # ٢) پاشەکەوت: SRT + force_style (هەر بە ناوی ڕاستەقینەی فۆنت)
        fam = _ass_font_family(font_path) if font_path else "Amiri"
        vf = "subtitles=burn_subs.srt:fontsdir=%s:force_style='%s'" % (
            _ff_quote(fontsdir), _ass_style(fam))
    out_tmp = os.path.join(cwd, "burn_out.mp4")
    cmd = ["ffmpeg", "-y", "-loglevel", "error", "-i", src_local]
    if voice_track and os.path.exists(voice_track):
        # دەنگی کوردی جیگای دەنگی ڕەسەن دەگرێتەوە
        cmd += ["-i", os.path.abspath(voice_track),
                "-map", "0:v:0", "-map", "1:a:0", "-shortest"]
    cmd += ["-vf", vf,
            "-c:v", "libx264", "-preset", "veryfast", "-crf", "23",
            "-c:a", "aac", "-b:a", "160k",
            "-movflags", "+faststart", out_tmp]
    try:
        res = subprocess.run(cmd, capture_output=True, text=True,
                             timeout=timeout, cwd=cwd)
    except subprocess.TimeoutExpired:
        raise BotError(ERR_LONG)
    if res.returncode != 0 or not os.path.exists(out_tmp) or os.path.getsize(out_tmp) == 0:
        log.warning("burn failed: %s", (res.stderr or "")[-600:])
        return None
    shutil.move(out_tmp, out_path)
    return out_path


def tg_send_video(chat, path, caption=None):
    payload = {"chat_id": chat, "supports_streaming": "true"}
    if caption:
        payload["caption"] = caption[:1000]
    try:
        with open(path, "rb") as fh:
            tg_call("sendVideo", payload,
                    files={"video": (os.path.basename(path), fh)}, timeout=900)
        return True
    except Exception as e:
        log.warning("sendVideo failed (%s) — fallback to sendDocument", e)
        try:
            tg_send_doc(chat, path, caption)
            return True
        except Exception as e2:
            log.error("video send fallback failed: %s", e2)
            return False


_whisper_model = None
_HARD_DEADLINE = None   # کۆتایی کاتی خول — لە main() دادەنرێت (بۆ ڕاوەستاندنی ئەرکی قورس)


def get_whisper():
    global _whisper_model
    if _whisper_model is None:
        from faster_whisper import WhisperModel
        log.info("loading whisper model: %s", WHISPER_MODEL)
        _whisper_model = WhisperModel(WHISPER_MODEL, device="cpu", compute_type="int8")
    return _whisper_model


def transcribe_file(wav_path, t0):
    model = get_whisper()
    seg_iter, info = model.transcribe(
        wav_path,
        language=SOURCE_LANG,
        vad_filter=True,
        vad_parameters={"min_silence_duration_ms": 500},
    )
    segs = []
    for s in seg_iter:
        if time.time() - t0 > HEAVY_BUDGET:
            raise BotError(ERR_LONG)
        if _HARD_DEADLINE and time.time() > _HARD_DEADLINE:
            raise BotError(ERR_LONG)
        txt = (s.text or "").strip()
        if txt:
            segs.append((float(s.start), float(s.end), txt))
    return segs, info


def merge_cues(segs, max_dur=12.0, max_gap=1.2):
    """یەکخستنی بەشە بچووکەکان بۆ ژێرنووسی خوێندنئاسان."""
    out, cur = [], None
    for st, en, txt in segs:
        if cur is not None and (en - cur[1]) <= max_dur and (st - cur[1]) <= max_gap:
            cur = (cur[0], en, (cur[2] + " " + txt).strip())
        else:
            if cur is not None:
                out.append(cur)
            cur = (st, en, txt)
    if cur is not None:
        out.append(cur)
    return out


def split_long_cues(cues, max_dur=12.0, max_chars=95):
    """بەشە زۆر درێژەکان دەکات بە چەند ژێرنووسی بچووکتر:
    - خوێندنئاسانتر لەسەر ڤیدیۆ
    - وەرگێڕان بە پارچەی بچووکتر — متمانەبەخۆتر و خێراتر
    کاتەکان بەپێی درێژی دەق بە شێوەی ڕێژەیی دابەش دەکرێن."""
    out = []
    for st, en, txt in cues:
        txt = (txt or "").strip()
        if not txt:
            continue
        dur = max(0.4, float(en) - float(st))
        if dur <= max_dur and len(txt) <= max_chars:
            out.append((float(st), float(en), txt))
            continue
        words = txt.split()
        cps = len(txt) / dur          # پیت بۆ چرکە
        t_cur = float(st)
        buf = ""
        for w in words:
            cand = (buf + " " + w) if buf else w
            over_time = (t_cur + len(cand) / cps - float(st)) > dur
            if len(cand) > max_chars or over_time:
                if buf:
                    t_next = min(float(en), t_cur + len(buf) / cps)
                    out.append((t_cur, max(t_next, t_cur + 0.4), buf))
                    t_cur = t_next
                buf = w
            else:
                buf = cand
        if buf:
            out.append((t_cur, float(en), buf))
    return out


def render_srt(cues, texts):
    blocks = []
    for i, ((st, en, _), t) in enumerate(zip(cues, texts), 1):
        blocks.append("%d\n%s --> %s\n%s\n" % (i, ts(st), ts(en), (t or "").strip()))
    return "\n".join(blocks)


def build_transcript_docx(cues, texts, out_path, title=""):
    doc = DocxDocument()
    if title:
        p = doc.add_paragraph(title[:120])
        _set_bidi_docx(p)
    for (st, en, _), t in zip(cues, texts):
        p = doc.add_paragraph("[%s] %s" % (ts(st), (t or "").strip()))
        _set_bidi_docx(p)
    doc.save(out_path)


def ytdl_download(url, workdir):
    import yt_dlp
    import yt_dlp.utils

    opts = {
        "format": ("best[height<=720]+bestaudio/best[height<=720]/bestaudio/best"
                   if BURN_SUBTITLES else "bestaudio/best"),  # ڤیدیۆ پێویستە بۆ ژێرنووس
        "outtmpl": os.path.join(workdir, "src.%(ext)s"),
        "noplaylist": True,
        "quiet": True,
        "no_warnings": True,
        "socket_timeout": 30,
        "retries": 3,
        "restrictfilenames": True,
    }
    if MAX_VIDEO_MINUTES > 0:
        opts["match_filter"] = yt_dlp.utils.match_filter_func(
            "duration<=%d" % (MAX_VIDEO_MINUTES * 60))
    if YT_COOKIES:
        cf = os.path.join(workdir, "cookies.txt")
        with open(cf, "w", encoding="utf-8") as fh:
            fh.write(YT_COOKIES + "\n")
        opts["cookiefile"] = cf

    with yt_dlp.YoutubeDL(opts) as ydl:
        info = ydl.extract_info(url, download=True)
        if info.get("_type") == "playlist" and info.get("entries"):
            info = info["entries"][0]
        path = None
        for d in (info.get("requested_downloads") or []):
            if d.get("filepath"):
                path = d["filepath"]
                break
        if not path:
            path = ydl.prepare_filename(info)
        if not os.path.exists(path):
            g = sorted(glob.glob(os.path.join(workdir, "src.*")))
            if not g:
                raise RuntimeError("downloaded file not found")
            path = g[-1]
        title = info.get("title") or "video"
        duration = info.get("duration") or 0
    return path, title, duration


def run_speech_pipeline(src_path, chat, workdir, state=None):
    t0 = time.time()
    tg_action(chat, "typing")
    # ---- ئەگەر ڤیدیۆ زۆر گەورەیە، پێش هەموو شتێک بچووکی دەکەینەوە ----
    src_path = shrink_source_for_processing(src_path, workdir)
    wav = os.path.join(workdir, "audio_16k.wav")
    ffmpeg_extract(src_path, wav)
    tg_send(chat, "دەنگەکە دەخوێنمەوە... (ئەم قۆناغە ماوەیەک دەخایەنێت)")
    segs, info = transcribe_file(wav, t0)
    if not segs:
        raise BotError(ERR_AUDIO)
    lang = getattr(info, "language", None) or "?"
    dur = getattr(info, "duration", 0) or (segs[-1][1] if segs else 0)
    cues = split_long_cues(merge_cues(segs))
    tg_send(chat, MSG_TRANSCRIBED.format(lang=lang, dur=fmt_dur(dur), n=len(cues)))
    texts = [c[2] for c in cues]
    outs, nfail = translate_map_safe(texts)
    if nfail >= len(cues):
        # هەموو وەرگێڕانەکە سەرنەچوو — لانیکەم نووسینی ڕەسەن بنێرەوە
        try:
            orig = os.path.join(workdir, "transcript_original.txt")
            with open(orig, "w", encoding="utf-8") as fh:
                fh.write("\n".join(texts))
            _safe_send(chat, "هەڵە: " + ERR_TRANS)
            tg_send_doc(chat, orig, MSG_ORIG_TRANSCRIPT)
        except Exception:
            pass
        return
    if nfail:
        _safe_send(chat, MSG_TRANS_PARTIAL.format(n=nfail))
    tg_action(chat, "upload_document")
    base = sanitize(os.path.splitext(os.path.basename(src_path))[0]) or "video"
    srt_path = os.path.join(workdir, "sorani_%s.srt" % base)
    with open(srt_path, "w", encoding="utf-8") as fh:
        fh.write(render_srt(cues, outs))
    docx_path = os.path.join(workdir, "sorani_%s.docx" % base)
    build_transcript_docx(cues, outs, docx_path, title="وەرگێڕانی کوردی سۆرانی — %s" % base)
    pdf_path = None
    if SEND_PDF_COPY:
        try:
            blocks = blocks_from_transcript(cues, outs)
            pdf_path = build_pdf_blocks(
                blocks, os.path.join(workdir, "sorani_%s.pdf" % base),
                title="وەرگێڕانی کوردی سۆرانی — %s" % base)
        except Exception as e:
            log.warning("transcript PDF failed: %s", e)
    preview = "\n".join(x for x in outs if x.strip())[:700]
    tg_send(chat, preview + ("…" if len(preview) == 700 else ""))
    tg_send_doc(chat, srt_path, MSG_VIDEO_DONE)
    tg_send_doc(chat, docx_path, None)
    if pdf_path:
        tg_send_doc(chat, pdf_path, None)
    # ---- دەنگی کوردی سۆرانی (TTS) ئەگەر بەکارهێنەر چالاکی کردبێت ----
    voice_on = bool((((state or {}).get("voice")) or {}).get(str(chat)))
    voice_track = None
    if voice_on:
        if not tts_available():
            _safe_send(chat, MSG_TTS_NOKEY)
        elif ffprobe_has_video(src_path):
            v_chk = float(dur or 0) or ffprobe_duration(src_path)
            if not (v_chk and v_chk > BURN_MAX_MINUTES * 60):
                tg_send(chat, MSG_TTS_START)
                tg_action(chat, "upload_video")
                try:
                    voice_track = build_kurdish_voice_track(cues, outs, dur, workdir)
                except BotError as e:
                    _safe_send(chat, "هەڵە: %s" % e)
                    voice_track = None
                if not voice_track:
                    _safe_send(chat, MSG_TTS_FAIL)
    # ---- ژێرنووسی کوردی لەسەر خۆی ڤیدیۆ (burn-in) ----
    if BURN_SUBTITLES:
        v_dur = float(dur or 0) or ffprobe_duration(src_path)
        if not ffprobe_has_video(src_path):
            log.info("burn skipped: source has no video stream (audio only)")
        elif v_dur and v_dur > BURN_MAX_MINUTES * 60:
            _safe_send(chat, MSG_BURN_SKIPPED.format(max=BURN_MAX_MINUTES))
        else:
            if not voice_track:
                tg_send(chat, MSG_BURN_START)
            tg_action(chat, "upload_video")
            vid_out = os.path.join(workdir, "sorani_%s_subtitled.mp4" % base)
            try:
                vid_path = burn_subtitles(src_path, srt_path, vid_out, workdir,
                                          voice_track=voice_track)
            except BotError as e:
                _safe_send(chat, "هەڵە: %s" % e)
                vid_path = None
            if vid_path and os.path.exists(vid_path):
                # ---- بچووککردنەوە بۆ خوار OUT_VIDEO_MAX_MB (بنەڕەت ٢٠ مێگابایت) ----
                limit_mb = OUT_VIDEO_MAX_MB
                if os.path.getsize(vid_path) > limit_mb * 1024 * 1024:
                    _safe_send(chat, MSG_COMPRESS.format(mb=limit_mb))
                    tg_action(chat, "upload_video")
                    small = os.path.join(workdir, "sorani_%s_subtitled_small.mp4" % base)
                    try:
                        cpath = compress_video(vid_path, small,
                                               max(8.0, limit_mb - 1.0), workdir)
                    except Exception as e:
                        log.warning("compress failed: %s", e)
                        cpath = None
                    if (cpath and os.path.exists(cpath) and os.path.getsize(cpath) > 0
                            and os.path.getsize(cpath) < os.path.getsize(vid_path)):
                        vid_path = cpath
                if os.path.getsize(vid_path) <= 49 * 1024 * 1024:
                    tg_send_video(chat, vid_path, CAPTION_BURN)
                else:
                    _safe_send(chat, MSG_BURN_FAIL)
            else:
                _safe_send(chat, MSG_BURN_FAIL)


# --------------------------- هاندەرەکان -------------------------------

def route_document(name, mime):
    ext = os.path.splitext(name or "")[1].lower()
    if ext in DOC_EXTS:
        return ext
    if mime == "application/pdf":
        return ".pdf"
    if mime == "application/x-subrip":
        return ".srt"
    if mime == "text/plain":
        return ".txt"
    if mime and "wordprocessingml" in mime:
        return ".docx"
    if mime and "presentationml" in mime:
        return ".pptx"
    return None


def route_media(name, mime):
    ext = os.path.splitext(name or "")[1].lower()
    if ext in MEDIA_EXTS:
        return ext
    if mime and (mime.startswith("video/") or mime.startswith("audio/")):
        return ext or ".bin"
    return None


def handle_document_file(tg_file, name, chat, workdir):
    kind = route_document(name, tg_file.get("mime_type") or "")
    if not kind:
        raise BotError(MSG_UNSUPPORTED)
    if int(tg_file.get("file_size") or 0) > TG_MAX_FILE:
        raise BotError(err_big())
    tg_send(chat, MSG_DOC_START)
    path = tg_download(tg_file["file_id"], workdir, name)
    with open(path, "rb") as fh:
        data = fh.read()
    stem = sanitize(os.path.splitext(os.path.basename(name))[0]) or "file"
    out_path = None
    pdf_path = None
    caption = CAPTION_DOC

    if kind == ".pdf":
        pages = extract_pdf_text_or_ocr(data, workdir)
        tpages, nfail = translate_map_safe(pages, workers=4)
        if nfail >= len(pages):
            raise BotError(ERR_TRANS)
        try:
            out_path = build_pdf(tpages,
                                 os.path.join(workdir, "sorani_%s.pdf" % stem),
                                 title=stem)
        except Exception as e:
            log.warning("PDF build failed, falling back to TXT: %s", e)
            out_path = os.path.join(workdir, "sorani_%s.txt" % stem)
            with open(out_path, "w", encoding="utf-8") as fh:
                fh.write("\n\n".join(tpages))
            caption = CAPTION_DOC + " (تێبینی: دروستکردنی PDF سەرنەچوو، بە TXT نێردرا)"

    elif kind == ".docx":
        out_path = os.path.join(workdir, "sorani_%s.docx" % stem)
        with open(out_path, "wb") as fh:
            fh.write(process_docx(data))
        if SEND_PDF_COPY:
            try:
                ensure_font_and_register(workdir)
                pdf_path = docx_to_pdf(out_path, workdir, title=stem)
            except Exception as e:
                log.warning("docx→pdf failed: %s", e)

    elif kind == ".pptx":
        out_path = os.path.join(workdir, "sorani_%s.pptx" % stem)
        with open(out_path, "wb") as fh:
            fh.write(process_pptx(data))
        if SEND_PDF_COPY:
            try:
                ensure_font_and_register(workdir)
                pdf_path = pptx_to_pdf(out_path, workdir, title=stem)
            except Exception as e:
                log.warning("pptx→pdf failed: %s", e)

    elif kind == ".txt":
        out_path = os.path.join(workdir, "sorani_%s.txt" % stem)
        with open(out_path, "wb") as fh:
            fh.write(process_txt(data))
        if SEND_PDF_COPY:
            try:
                with open(out_path, "r", encoding="utf-8") as fh:
                    ttext = fh.read()
                blocks = blocks_from_plain(ttext)
                if blocks:
                    pdf_path = build_pdf_blocks(
                        blocks, os.path.join(workdir, "sorani_%s.pdf" % stem),
                        title=stem)
            except Exception as e:
                log.warning("txt→pdf failed: %s", e)

    elif kind == ".srt":
        srt_out, scues, souts = translate_srt(data)
        if not srt_out:
            raise BotError(ERR_GEN)
        out_path = os.path.join(workdir, "sorani_%s.srt" % stem)
        with open(out_path, "w", encoding="utf-8") as fh:
            fh.write(srt_out)
        if SEND_PDF_COPY and scues:
            try:
                blocks = blocks_from_transcript(scues, souts)
                pdf_path = build_pdf_blocks(
                    blocks, os.path.join(workdir, "sorani_%s.pdf" % stem),
                    title=stem)
            except Exception as e:
                log.warning("srt→pdf failed: %s", e)

    if out_path and os.path.exists(out_path):
        tg_send_doc(chat, out_path, caption)
    else:
        raise BotError(ERR_GEN)
    if pdf_path and os.path.exists(pdf_path):
        tg_send_doc(chat, pdf_path, CAPTION_PDF)


def handle_media_file(tg_file, chat, workdir, name_hint="media", state=None):
    if int(tg_file.get("file_size") or 0) > TG_MAX_FILE:
        raise BotError(err_big())
    tg_send(chat, MSG_MEDIA_START)
    path = tg_download(tg_file["file_id"], workdir, name_hint)
    run_speech_pipeline(path, chat, workdir, state)


def handle_link(url, chat, workdir, state=None):
    tg_action(chat, "typing")
    tg_send(chat, PROG_DOWNLOAD)
    try:
        src_path, title, duration = ytdl_download(url, workdir)
    except BotError:
        raise
    except Exception as e:
        log.warning("yt-dlp failed for %s: %s", url, e)
        raise BotError(ERR_LINK)
    msg = "داگیرا: %s" % title[:120]
    if duration:
        msg += "\nماوە: %s" % fmt_dur(duration)
    tg_send(chat, msg)
    run_speech_pipeline(src_path, chat, workdir, state)


def handle_text_translate(text, chat):
    tg_action(chat, "typing")
    out = translate_text(text)
    if out.strip() == text.strip():
        tg_send(chat, MSG_ALREADY)
    else:
        tg_send(chat, out)


def handle_command(text, chat, state):
    """فەرمانەکان — True ئەگەر مامەڵە کرابێت."""
    parts = text.split()
    cmd = (parts[0] if parts else "").split("@")[0].lower()
    is_dict = isinstance(state, dict)
    voice = state.setdefault("voice", {}) if is_dict else {}
    if cmd in ("/voice_on", "/deng_on", "/voiceon", "/dengon"):
        if is_dict:
            voice[str(chat)] = True
            save_state(state)
        tg_send(chat, MSG_VOICE_ON if tts_available() else MSG_TTS_NOKEY)
        return True
    if cmd in ("/voice_off", "/deng_off", "/voiceoff", "/dengoff"):
        if is_dict:
            voice[str(chat)] = False
            save_state(state)
        tg_send(chat, MSG_VOICE_OFF)
        return True
    if cmd in ("/voice", "/deng", "/voice_status", "/settings"):
        on = voice.get(str(chat)) if is_dict else False
        tg_send(chat, MSG_VOICE_STATUS_ON if on else MSG_VOICE_STATUS_OFF)
        return True
    if cmd.startswith("/"):
        tg_send(chat, WELCOME)
        return True
    return False


def handle_message(msg, chat, workdir, state=None):
    doc = msg.get("document")
    if doc:
        name = doc.get("file_name") or "file"
        if route_document(name, doc.get("mime_type") or ""):
            handle_document_file(doc, name, chat, workdir)
        elif route_media(name, doc.get("mime_type") or ""):
            handle_media_file(doc, chat, workdir, os.path.basename(name), state)
        else:
            raise BotError(MSG_UNSUPPORTED)
        return

    for key in ("video", "voice", "audio", "video_note"):
        if key in msg:
            handle_media_file(msg[key], chat, workdir, key, state)
            return

    text = (msg.get("text") or "").strip()
    if not text:
        text = (msg.get("caption") or "").strip()
    if not text:
        raise BotError(MSG_UNSUPPORTED)

    if text.startswith("/"):
        handle_command(text, chat, state)
        return

    m = re.search(r"https?://\S+", text)
    if m:
        handle_link(m.group(0).rstrip(".,;)]}>"), chat, workdir, state)
        return

    handle_text_translate(text, chat)


def process_update(u, workdir, state=None):
    msg = u.get("message") or {}
    chat = (msg.get("chat") or {}).get("id")
    if not chat:
        return
    task_dir = tempfile.mkdtemp(prefix="task_", dir=workdir)
    try:
        handle_message(msg, chat, task_dir, state)
    except BotError as e:
        _safe_send(chat, "هەڵە: %s" % e)
    except Exception:
        log.error("update failed:\n%s", traceback.format_exc())
        _safe_send(chat, ERR_GEN)
    finally:
        shutil.rmtree(task_dir, ignore_errors=True)


# ------------------------------ خولی سەرەکی ---------------------------

def self_check():
    print("== sorani-translator-bot :: self check ==")
    print("BOT_TOKEN:", "set" if BOT_TOKEN else "MISSING (add GitHub secret)")
    print("GLM fallback:", ("enabled, models=" + ",".join(GLM_MODELS)) if GLM_API_KEY else "not configured (Google only)")
    print("whisper model:", WHISPER_MODEL, "| max video minutes:", MAX_VIDEO_MINUTES or "unlimited")
    print("ffmpeg:", shutil.which("ffmpeg") or "MISSING")
    print("libreoffice:", shutil.which("soffice") or "missing (ReportLab fallback will be used)")
    print("tesseract OCR:", shutil.which("tesseract") or "missing (scanned PDFs will fail)")
    print("burn subtitles:", ("on, max %d min" % BURN_MAX_MINUTES) if BURN_SUBTITLES else "off")
    print("pdf copy:", "on" if SEND_PDF_COPY else "off")
    print("output video limit:", "%d MB (auto-compress)" % OUT_VIDEO_MAX_MB)
    print("telegram api:", ("LOCAL Bot API server (%s) — files up to 2GB" % TG_BASE)
          if LOCAL_BOT_API else "official api.telegram.org — downloads up to 20MB")
    print("kurdish tts:", ("on, speaker=%s" % TTS_SPEAKER) if TTS_API_KEY
          else "off (set KURDISH_TTS_API_KEY secret)")
    try:
        import yt_dlp
        print("yt-dlp:", yt_dlp.version.__version__)
    except Exception as e:
        print("yt-dlp: FAIL", e)
    try:
        import faster_whisper  # noqa
        print("faster-whisper: OK")
    except Exception as e:
        print("faster-whisper: FAIL", e)
    for i, (host, client, use_dj) in enumerate(GTX_HOSTS, 1):
        name = host.split("/")[2]
        try:
            out, src = _gtx_request(host, client, use_dj, "Hello, how are you today?", False)
            print("Google route %d (%s): OK (src=%s) sample: %s" % (i, name, src, out[:60]))
        except Exception as e:
            print("Google route %d (%s): FAIL %s" % (i, name, e))
    if TTS_API_KEY:
        try:
            b = _tts_once("سڵاو")
            print("Kurdish TTS: OK (%d bytes)" % len(b))
        except Exception as e:
            print("Kurdish TTS: FAIL", e)
    print("done.")


def main():
    global _HARD_DEADLINE
    if "--check" in sys.argv:
        self_check()
        return
    if not BOT_TOKEN:
        print("BOT_TOKEN is missing — add it as a GitHub secret first.")
        sys.exit(42)

    state = load_state()
    save_state(state)

    workdir = tempfile.mkdtemp(prefix="soranibot_")
    start = time.time()
    deadline = start + RUN_SECONDS
    _HARD_DEADLINE = deadline - 600   # ١٠ خولەک بۆ ناردنەوەی دەرئەنجامەکان حیجز بکە

    if state.get("offset", 0) <= 0:
        try:
            skip_backlog(state)
        except Exception as e:
            log.warning("skip_backlog failed: %s", e)

    log.info("bot cycle started (offset=%s, budget=%ss)", state.get("offset"), RUN_SECONDS)

    while time.time() < deadline:
        remaining = deadline - time.time()
        updates = tg_get_updates(state.get("offset", 0),
                                 timeout=int(min(45, max(5, remaining))))
        if (not updates and IDLE_EXIT_SECONDS > 0
                and (time.time() - start) > IDLE_EXIT_SECONDS):
            break
        for u in updates:
            try:
                process_update(u, workdir, state)
            except Exception:
                log.error("process_update crashed:\n%s", traceback.format_exc())
            state["offset"] = u["update_id"] + 1
            save_state(state)
            log.info("processed update %s", u["update_id"])

    save_state(state)
    shutil.rmtree(workdir, ignore_errors=True)
    log.info("bot cycle finished")


if __name__ == "__main__":
    main()
